"""Séries por temperatura, curvas e exportação (ESPECIFICACAO.md, seção 6)."""

from __future__ import annotations

import io

import pytest

from tests import synthetic as syn
from tests.conftest import COORD, PESQ, confirm, upload


def _rock(analysis, sample, toc, curves=True):
    return {"analysis": analysis, "sample": sample, "toc": toc, "tmax": 440, "s1": 1.0, "s2": 10.0, "hi": 300, "oi": 10, "curves": curves}


@pytest.fixture
def loaded(client, project):
    """Réplicas de experimento A/B a 300 °C (rocha H e E), SE a 300 °C e
    rocha original — tudo sintético."""
    htm = syn.rockeval_htm(
        [
            _rock("r_HP300NA_1", "HP300NA", 10.0),
            _rock("r_HP300NA_2", "HP300NA", 12.0),
            _rock("r_HP300NB_1", "HP300NB", 14.0),
            _rock("r_HP300NAE_1", "HP300NAE", 6.0),
            _rock("r_HP300SE_1", "HP300SE", 20.0),
            _rock("r_ROA_1", "ROA", 30.0, curves=False),
        ]
    )
    pv = upload(client, project["id"], {"job.htm": htm})
    # HP300NA/HP300NB como rocha: fração "HP" (sem sufixo) → trata como H
    decisions = [
        {"row": r["row"], "action": "create", "code": r["name"], "fraction": "H"} for r in pv["rows"] if r["name"] in ("HP300NA", "HP300NB")
    ]
    confirm(client, project["id"], pv, decisions)
    return project


def _series(client, pid, **params):
    query = "&".join(f"{k}={v}" for k, v in params.items())
    r = client.get(f"projects/{pid}/series?{query}", headers=PESQ)
    assert r.status_code == 200, r.text
    return r.json()


def test_series_aggregates_experiment_replicates_and_groups_se_with_h(client, loaded):
    data = _series(client, loaded["id"], technique="rockeval", parameter="TOC")
    by = {s["fraction"]: s for s in data["series"]}
    assert set(by) == {"H", "E"}
    h_point = by["H"]["points"][0]
    # amostras HP300NA (média 11), HP300NB (14) e HP300SE (20) → média das médias
    assert h_point["temperature_c"] == 300 and h_point["n_samples"] == 3
    assert h_point["mean"] == pytest.approx((11 + 14 + 20) / 3)
    assert "SE" in by["H"]["fractions"]
    assert by["E"]["points"][0]["mean"] == 6.0
    assert [b["code"] for b in data["baseline"]] == ["ROA"]
    assert data["unit"] == "%"


def test_series_split_by_experiment_replicate(client, loaded):
    data = _series(client, loaded["id"], technique="rockeval", parameter="TOC", split_replicates="true")
    labels = {s["label"]: s["points"][0]["mean"] for s in data["series"] if s["fraction"] == "H"}
    assert any(lbl.endswith("réplica A") and v == 11.0 for lbl, v in labels.items())
    assert any(lbl.endswith("réplica B") and v == 14.0 for lbl, v in labels.items())


def test_validity_modes(client, loaded):
    pid = loaded["id"]
    table = {s["code"]: s for s in client.get(f"projects/{pid}/samples", headers=PESQ).json()["samples"]}
    client.post(f"projects/{pid}/samples/{table['HP300SE']['id']}/validation", json={"valid": False}, headers=COORD)
    client.post(f"projects/{pid}/samples/{table['HP300NA']['id']}/validation", json={"valid": True}, headers=COORD)
    default = _series(client, pid, technique="rockeval", parameter="TOC")
    h = next(s for s in default["series"] if s["fraction"] == "H")["points"][0]
    assert h["n_samples"] == 2  # SE inválida fica de fora
    only_valid = _series(client, pid, technique="rockeval", parameter="TOC", mode="validas")
    h = next(s for s in only_valid["series"] if s["fraction"] == "H")["points"][0]
    assert h["n_samples"] == 1 and h["mean"] == 11.0 and h["sd"] == pytest.approx(1.4142, rel=1e-3)
    everything = _series(client, pid, technique="rockeval", parameter="TOC", mode="todas")
    assert next(s for s in everything["series"] if s["fraction"] == "H")["points"][0]["n_samples"] == 3


def test_pyrogram_curves(client, loaded):
    pid = loaded["id"]
    table = {s["code"]: s for s in client.get(f"projects/{pid}/samples", headers=PESQ).json()["samples"]}
    ids = f"{table['HP300NA']['id']},{table['ROA']['id']}"
    curves = client.get(f"projects/{pid}/analysis-data?technique=rockeval&sample_ids={ids}", headers=PESQ).json()
    assert [c["sample_code"] for c in curves] == ["HP300NA", "HP300NA"]  # ROA sem curvas
    assert "HC" in curves[0]["data"]["pyro"]["series"]


def test_export_csv_and_xlsx(client, loaded):
    pid = loaded["id"]
    r = client.post(f"projects/{pid}/export", json={"format": "csv", "columns": ["rockeval.TOC"]}, headers=PESQ)
    assert r.status_code == 200 and "text/csv" in r.headers["content-type"]
    text = r.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Amostra;Fração;Temperatura")
    assert "Rock-Eval COT (%) média" in text and "HP300NA" in text
    r = client.post(f"projects/{pid}/export", json={"format": "xlsx", "sample_ids": [], "columns": []}, headers=PESQ)
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.active.cell(row=1, column=1).value == "Amostra"
    assert wb.active.max_row == 6  # cabeçalho + 5 amostras


def test_catalog_lists_techniques(client):
    keys = [t["key"] for t in client.get("catalog", headers=PESQ).json()]
    assert keys == ["chnso", "leco", "leco_ri", "rockeval", "gc_fid", "gc_tcd", "gas_balanco", "pygcms"]
