"""Estudo do PostRun (RE7S) → projeto de polímeros (22/10/2026). Planilha
SINTÉTICA montada aqui com openpyxl, no formato versão 1 do contrato."""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.parsers import ParseError, parse_file
from app.services import profiles
from app.services.catalog import param, postrun_param
from app.services.codes import parse_polymer_code
from tests.conftest import COORD, PESQ, confirm, upload

HEADER = [
    "Estudo", "Método", "Amostra", "Réplica", "Composição", "Componente", "Parâmetro", "Valor", "Unidade", "U95",
    "Fração nominal (%)", "Fração real (%)", "Massa da amostra (mg)", "Data",
]  # fmt: skip


P50, P75, PS = "PP 50 % · PS 50 %", "PP 75 % · PS 25 %", "PS 100 %"


def _rows(study="Estudo X", methods=("gauss", "molde"), shift=0.0):
    rows = []
    runs = [
        # amostra, réplica, composição, {componente: (nominal, real, recuperação)}, massa
        # "Amostra" = chave do grupo de réplicas = o texto da composição nominal (contrato v1)
        (P50, "2026-01-10_PP_PS_50-50_1", P50, {"PP": (50.0, 49.5, 98.0), "PS": (50.0, 50.5, 101.0)}, 2.0),
        (P50, "2026-01-10_PP_PS_50-50_2", P50, {"PP": (50.0, 50.2, 96.0), "PS": (50.0, 49.8, 99.0)}, 2.1),
        (P75, "2026-01-11_PP_PS_75-25_1", P75, {"PP": (75.0, None, 97.0), "PS": (25.0, None, 103.0)}, 2.0),  # sem fração real
        (PS, "2026-01-12_PS_1", PS, {"PS": (100.0, 100.0, 100.0)}, 1.5),
    ]
    for method in methods:
        delta = 0.0 if method == "gauss" else -1.0
        for sample, rep, comp, parts, mass in runs:
            for c, (nom, real, rec) in parts.items():
                common = [study, method, sample, rep, comp, c]
                tail = [nom, real, mass, "2026-01-10"]
                rows.append(common + ["Recuperação", rec + delta + shift, "%", 1.5] + tail)
                frac = real if real is not None else nom
                rows.append(common + ["Massa encontrada", mass * frac / 100 * rec / 100, "mg", 0.02] + tail)
                rows.append(common + ["Massa esperada", mass * frac / 100, "mg", None] + tail)
                rows.append(common + ["ΔTmax", 1.2, "°C", None] + tail)
                rows.append(common + ["Δlargura", -3.0, "%", None] + tail)
            run = [study, method, sample, rep, comp, None]
            rows.append(run + ["Recuperação total", 99.0 + delta + shift, "%", 2.0, None, None, mass, "2026-01-10"])
            rows.append(run + ["Massa encontrada total", mass * 0.99, "mg", None, None, None, mass, "2026-01-10"])
            rows.append(run + ["R²", 0.998, None, None, None, None, mass, "2026-01-10"])
    return rows


def make_study(fmt="horun-postrun-estudo", metodo="molde", status="finalizado", gerado="2026-01-20T10:00:00", **kw) -> bytes:
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = "Sobre"
    about.append(["Chave", "Valor"])
    for k, v in [("formato", fmt), ("versao", 1), ("estudo", kw.get("study", "Estudo X")), ("descricao", "sintético"), ("metodo", metodo),
                 ("canal", "pyro:FID"), ("gerado_em", gerado), ("status", status)]:  # fmt: skip
        about.append([k, v])
    wb.create_sheet("Análises").append(["qualquer coisa", 1, 2, 3])  # abas para gente — ignoradas
    ws = wb.create_sheet("Horun_Resultados")
    ws.append(HEADER)
    for r in _rows(**kw):
        ws.append(r)
    wb.create_sheet("Médias").append(["Amostra", "Recuperação"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _vals(record) -> dict[str, float]:
    return {v.parameter: v.value for v in record.values}


def test_parser_maps_replicates_components_methods_and_variables():
    res = parse_file("estudo.xlsx", make_study())
    assert res.technique == "postrun_estudo" and len(res.records) == 4
    rec = next(r for r in res.records if r.raw_name == "2026-01-10_PP_PS_50-50_2")
    assert rec.code_hint == P50 and rec.key == "Estudo X|2026-01-10_PP_PS_50-50_2"
    assert rec.extra["replicate_hint"] == 2 and rec.analyzed_at == "2026-01-10"
    v = _vals(rec)
    # método principal = molde (Sobre): chave sem sufixo = molde; os dois métodos com sufixo
    assert v["rec_PP"] == 95.0 and v["rec_PP__molde"] == 95.0 and v["rec_PP__gauss"] == 96.0
    assert v["rec_PP__u95"] == 1.5 and "rec_PP__u95" in v and "mexp_PP__u95" not in v
    assert v["rec_total"] == 98.0 and v["r2"] == 0.998 and v["sample_mass"] == 2.1
    assert v["frac_real_PP"] == 50.2 and v["dtmax_PS"] == 1.2 and v["dwidth_PS"] == -3.0
    sv = {x["label"]: x["value"] for x in rec.extra["sample_variables"]}
    assert sv == {"Fração PP": 50.0, "Fração PS": 50.0, "Composição": "PP 50 % · PS 50 %"}
    assert "frac_real_PP" not in _vals(next(r for r in res.records if r.code_hint == P75))  # fração real vazia
    pure = next(r for r in res.records if r.code_hint == PS)
    # puro: o componente do estudo que falta na amostra entra como 0 %
    assert {x["label"]: x["value"] for x in pure.extra["sample_variables"]}["Fração PP"] == 0.0
    assert not res.warnings


def test_parser_rejects_other_format_and_warns_open_study():
    with pytest.raises(ParseError, match="não é um Estudo do PostRun"):
        parse_file("x.xlsx", make_study(fmt="outra-coisa"))
    res = parse_file("x.xlsx", make_study(status="aberto", metodo="", methods=("gauss",)))
    assert any("ABERTO" in w for w in res.warnings)
    assert _vals(res.records[0])["rec_PP"] == _vals(res.records[0])["rec_PP__gauss"]  # sem "metodo": o único das linhas


def test_dynamic_catalog_params():
    p = postrun_param("rec_PEAD__molde")
    assert p and p.label == "Recuperação PEAD (moldes)" and p.unit == "%" and not p.main
    assert postrun_param("frac_real_PP").label == "Fração real PP"
    assert postrun_param("rec_PP").main and postrun_param("nada") is None
    assert param("postrun_estudo", "rec_total").main and param("postrun_estudo", "mfound_PS__u95").label == "U95 Massa encontrada PS"


@pytest.mark.parametrize(
    ("raw", "base", "rep"),
    [(P50, P50, None), ("PP 100 %", "PP 100 %", None), ("PP_PS_50-50", "PP_PS_50-50", None), ("PEAD_040", "PEAD_040", None), ("2026-01-10_PP_PS_50-50_7", "PP_PS_50-50", 7),
     ("2026-01-10_PEAD_040_2.B00", "PEAD_040", 2)],
)  # fmt: skip
def test_polymer_code_reader(raw, base, rep):
    info = parse_polymer_code(raw)
    assert (info.base_code, info.replicate, info.fraction) == (base, rep, "POL")


def test_polymer_profile_is_ready():
    assert profiles.validate({"sample_type": "polimeros_pirolise"}).resolved()["recipes"] == ["postrun_estudo", "rockeval"]


def _project(client):
    r = client.post("projects", json={"name": "Polímeros sintético", "profile": {"sample_type": "polimeros_pirolise"}}, headers=COORD)
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["profile"]["temperature_axis"] is False
    from tests.conftest import add_members

    add_members(client, p["id"], PESQ)
    return p


def test_import_groups_replicates_sets_variables_and_is_idempotent(client):
    p = _project(client)
    pid = p["id"]
    prev = upload(client, pid, {"estudo.xlsx": make_study()})
    assert prev["counts"]["ok"] == 1
    rows = {r["name"]: r for r in prev["rows"]}
    assert set(rows) == {P50, P75, PS}
    assert rows[P50]["n_measurements"] == 2 and rows[P50]["replicates"] == [1, 2]
    assert rows[PS]["suggested"]["fraction"] == "POL" and rows[PS]["suggested"]["code"] == PS
    res = confirm(client, pid, prev)
    assert res["samples_created"] == 3 and res["analyses_created"] == 4 and res["variables_set"] > 0

    variables = {v["label"]: v for v in client.get(f"projects/{pid}/variables", headers=PESQ).json()["variables"]}
    assert variables["Fração PP"]["kind"] == "numero" and variables["Fração PP"]["unit"] == "%"
    assert {o["value"] for o in variables["Composição"]["options"]} == {P50, P75, PS}
    samples = {s["code"]: s for s in client.get(f"projects/{pid}/samples", headers=PESQ).json()["samples"]}
    s50 = samples[P50]
    assert s50["variables"][variables["Fração PP"]["key"]] == 50.0
    assert samples[PS]["variables"][variables["Fração PP"]["key"]] == 0.0
    assert s50["values"]["postrun_estudo.rec_PP"]["mean"] == pytest.approx(96.0) and s50["values"]["postrun_estudo.rec_PP"]["n"] == 2

    # catálogo traz os parâmetros por componente que existem no banco
    cat = {t["key"]: t for t in client.get("catalog", headers=PESQ).json()}
    keys = {x["key"]: x for x in cat["postrun_estudo"]["params"]}
    assert keys["rec_PP"]["label"] == "Recuperação PP" and keys["rec_PP"]["main"] and "rec_PS__gauss" in keys

    # Séries em função da fração de PP: puro PS em 0, misturas em 50 e 75
    data = client.get(f"projects/{pid}/series?technique=postrun_estudo&parameter=rec_PS&variable={variables['Fração PP']['key']}", headers=PESQ).json()
    assert data["label"] == "Recuperação PS" and data["unit"] == "%"
    (line,) = data["series"]
    assert [pt["x"] for pt in line["points"]] == [0.0, 50.0, 75.0]

    # reimportar o mesmo estudo (outra exportação, números novos) atualiza, não duplica
    prev2 = upload(client, pid, {"estudo-v2.xlsx": make_study(gerado="2026-01-21T09:00:00", shift=1.0)})
    assert all(r["updates_existing"] for r in prev2["rows"]) and all(r["action"] == "link" for r in prev2["rows"])
    res2 = confirm(client, pid, prev2)
    assert res2["analyses_created"] == 0 and res2["analyses_updated"] == 4 and res2["samples_created"] == 0
    samples = {s["code"]: s for s in client.get(f"projects/{pid}/samples", headers=PESQ).json()["samples"]}
    assert samples[P50]["values"]["postrun_estudo.rec_PP"]["mean"] == pytest.approx(97.0)
    variables2 = client.get(f"projects/{pid}/variables", headers=PESQ).json()["variables"]
    assert len(variables2) == len(variables)

    # exportação e relatório aceitam os parâmetros dinâmicos
    r = client.post(f"projects/{pid}/export", json={"format": "csv", "columns": ["postrun_estudo.rec_PP"]}, headers=PESQ)
    assert r.status_code == 200 and "Recuperação PP" in r.content.decode("utf-8")
    r = client.post(f"projects/{pid}/export/report", json={"techniques": ["postrun_estudo"]}, headers=PESQ)
    assert r.status_code == 200, r.text
