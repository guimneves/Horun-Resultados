"""Arquivos SINTÉTICOS com a mesma estrutura dos reais (ESPECIFICACAO.md,
seção 3). Valores e nomes inventados — nunca copiar dados reais para cá."""

from __future__ import annotations

import csv
import io
import json

HEADER = [
    "Page", "1", "of", "1", "EuroVector", "Elemental", "Analyser", "AutoRun", "Name", "Date", "And", "Time", "Of",
    "Analysis", "Analysed", "By", "Configuration", "Calibration", "Type", "Signed", "By", "Operator", "Group",
]  # fmt: skip


def _pdf(pages: list[list[str]]) -> bytes:
    """Uma palavra por linha (como o pypdf entrega os PDFs do EuroVector)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for tokens in pages:
        y = A4[1] - 20
        x = 20
        c.setFont("Helvetica", 5)
        for tok in tokens:
            c.drawString(x, y, tok)
            y -= 7
            if y < 15:
                y = A4[1] - 20
                x += 120
        c.showPage()
    c.save()
    return buf.getvalue()


def _head(run: str, page: int, total: int, printed: str = "01 Jan 2026 - 10:00:00") -> list[str]:
    head = [printed, *HEADER]
    head[2], head[4] = str(page), str(total)
    return [*head, *run.split(" "), "05", "Mar", "2026", "-", "09:30:00", "EVR", "CHNS", "Linear", "evr", "GRP", "1"]


def chnso_summary_pdf(
    rows: list[tuple], run: str = "000001- Teste Sintetico", per_page: int = 6, printed: str = "01 Jan 2026 - 10:00:00"
) -> bytes:
    """rows: (posição, tipo Byp/Std/Smp, nome, N, C, H, S, O, W) com None = "-"."""
    pages = []
    chunks = [rows[i : i + per_page] for i in range(0, len(rows), per_page)] or [[]]
    for n, chunk in enumerate(chunks, start=1):
        tokens = _head(run, n, len(chunks), printed)
        if n == 1:
            tokens += ["Results", "Summary", "for", "Element", "%", "#", "Type", "Name", "N", "%", "C", "%", "H", "%", "S", "%", "O", "%", "W", "(mg)"]
        else:
            tokens = tokens[:5]
        for pos, typ, name, *vals in chunk:
            tokens += [str(pos), typ]
            parts = name.split(" ")
            # nomes quebrados como no real: "HP300-" "1"
            for part in parts:
                if "-" in part and not part.endswith("-"):
                    a, b = part.rsplit("-", 1)
                    tokens += [a + "-", b]
                else:
                    tokens.append(part)
            tokens += ["-" if v is None else f"{v:.3f}" for v in vals]
        pages.append(tokens)
    return _pdf(pages)


def chnso_single_pdf(position: int, typ: str, name: str, element_pct: dict, weight: float = 1.0, run: str = "000001- Teste Sintetico") -> bytes:
    tokens = _head(run, 1, 1)
    tokens += ["Single", "Sample", "Result", "Sample", "ID", "Sample", "Name", *name.split(" "), "Sample", "Position", "#", str(position)]
    tokens += ["Type", typ, "Sample", "Weight", "(mg)", f"{weight:.3f}"]
    if typ == "Std":
        tokens += ["Std", "Values", "%", "Substance", name, "Nitrogen", "99.999", "Carbon", "99.999"]
    tokens += ["Element", "Element", "%", "Area(uV*sec)", "Ret.Time", "(sec)", "Start", "(sec)", "End", "(sec)"]
    for el in ("Nitrogen", "Carbon", "Hydrogen", "Sulphur", "Oxygen"):
        v = element_pct.get(el)
        tokens += [el, *(["-"] * 5 if v is None else [f"{v:.3f}", "1,000.0", "100", "50", "150"])]
    return _pdf([tokens])


def chnso_other_pdf(title: str) -> bytes:
    return _pdf([_head("000001- Teste Sintetico", 1, 1) + title.split(" ") + ["#", "Type", "Name"]])


def plain_pdf(text: str) -> bytes:
    return _pdf([text.split(" ")])


def leco_diagnostic_pdf() -> bytes:
    # uma letra por linha, como o PDF de diagnóstico real
    return _pdf([list("SC832DR_00000") + list("Configuration") + list("ParameterValue") + list("ProductSC832DR")])


# ---------------------------------------------------------------- LECO

REP_HEADER = (
    "Date,Time,Set ID,Analysis Date,Analysis Time,Name,Type,Description,Comments,Operator,Include/Exclude State,"
    "Replicate ID,Replicate Guid,Repetition,Recalculation Date,Product,Method,Sample Mass,Carbon,Sulfur,Instrument Name"
).split(",")
SUM_DATE_HEADER = (
    "Date,Time,Set ID,Analysis Date,Name,Type,Description,Include/Exclude State,Number of Replicates,Instrument Name,"
    "Method,Carbon Average,Carbon Std. Dev.,Carbon Average,Sulfur Average,Sulfur Std. Dev.,Sulfur Average"
).split(",")
SUM_NAME_HEADER = (
    "Name,Type,Description,Set ID,Analysis Date,Date,Time,Include/Exclude State,Number of Replicates,Instrument Name,"
    "Method,Carbon Average,Carbon Std. Dev.,Carbon Average,Sulfur Average,Sulfur Std. Dev.,Sulfur Average"
).split(",")


def leco_csv(sets: list[dict]) -> bytes:
    """sets: {"name", "type", "set_id", "reps": [(C, S, massa, estado)]}.
    Reproduz a mistura real: resumo no layout "Date,..." antes das réplicas,
    resumo no layout "Name,..." depois, SEM repetir os cabeçalhos."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(SUM_DATE_HEADER)
    first = True
    for s in sets:
        n = len(s["reps"])
        c_avg = sum(r[0] for r in s["reps"]) / n
        s_avg = sum(r[1] for r in s["reps"]) / n
        date_summary = ["5/3/2026", "10:00:00 AM", s["set_id"], "5/3/2026 10:00:00 AM", s["name"], s.get("type", "Sample"), "",
                        "None are Excluded", str(n), "SC000 1", "Metodo X", f"{c_avg:.3g} %", "0.01 %", f"{c_avg:.3g} ± 0.01 %",
                        f"{s_avg:.3g} %", "0.01 %", f"{s_avg:.3g} ± 0.01 %"]  # fmt: skip
        w.writerow(date_summary)
        if first:
            w.writerow(REP_HEADER)
        for i, (c, sv, mass, state) in enumerate(s["reps"], start=1):
            w.writerow(["5/3/2026", f"10:0{i}:00 AM", s["set_id"], f"5/3/2026 10:0{i}:00 AM", "0:03:00", s["name"], s.get("type", "Sample"),
                        "", "", "Equipamento", state, "", "", str(i), "", "SC000", "Metodo X", f"{mass} g", f"{c} %", f"{sv} %", "SC000 1"])  # fmt: skip
        if first:
            w.writerow(SUM_NAME_HEADER)
            first = False
        w.writerow([s["name"], s.get("type", "Sample"), "", s["set_id"], "5/3/2026 10:00:00 AM", "5/3/2026", "10:00:00 AM",
                    "None are Excluded", str(n), "SC000 1", "Metodo X", f"{c_avg:.3g} %", "0.01 %", "x", f"{s_avg:.3g} %", "0.01 %", "x"])  # fmt: skip
    return buf.getvalue().encode("utf-8")


# ---------------------------------------------------------------- Rock-Eval

RE_HEADER = ["Analysis", "Sample", "Date", "Quantity (mg)", "Method", "Cycle", "TpkS2 (°C)", "Tmax (°C)", "TOC (%)", "MINC (%)",
             "S1 (mg/g)", "S2 (mg/g)", "S3 (mg/g)", "S1 S (%)", "HI", "OI", "S index"]  # fmt: skip


def rockeval_htm(analyses: list[dict]) -> bytes:
    """analyses: {"analysis", "sample", "toc", "tmax", "s1", "s2", "hi", "oi", "curves": bool}."""
    header = ["\xa0" + h.replace(" ", "\xa0") + "\xa0" for h in RE_HEADER]
    rows = []
    for a in analyses:
        rows.append([a["analysis"], a["sample"], "03/05/26 - 10h00", "60.0", "BULK ROCK", "BASIC", "460", str(a["tmax"]), str(a["toc"]),
                     "0.50", str(a["s1"]), str(a["s2"]), "0.40", "0.010", str(a["hi"]), str(a["oi"]), "1"])  # fmt: skip
    parts = ["<!DOCTYPE html><html><head><script>", 'var reportTitle = "Job report / Sintetico";']
    parts.append("var dataTableHeader = " + json.dumps(header, ensure_ascii=False) + ";\r\n")
    parts.append("var dataTableArray = " + json.dumps(rows) + ";\r\n")
    parts.append("var pyroCurveLegend = ['Time', 'Temp', 'HC', 'CO', 'CO2', 'SO2', 'C1', 'C4', 'H2', 'H2O', 'T°'];\r\n")
    links = []
    for i, a in enumerate(analyses):
        if a.get("curves", True):
            n = 50
            data = [list(range(n)), [300 + 8 * t for t in range(n)], [float(t % 17) for t in range(n)], [0.1] * n, [0.2] * n, [0.0] * n,
                    None, None, None, None]  # fmt: skip
            parts.append(f"var pyroData{i} = {json.dumps(data)};\r\n")
            links.append(f'["{a["analysis"]}",pyroData{i},null]')
        else:
            links.append(f'["{a["analysis"]}",null,null]')
    parts.append("var curveData = [" + ",\r\n".join(links) + "];\r\n")
    parts.append("</script></head><body></body></html>")
    return ("﻿" + "".join(parts)).encode("utf-8")


# ---------------------------------------------------------------- planilhas


def _xlsx(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gc_xlsx(detector: str, cell_code: str, analytes: list[tuple]) -> bytes:
    """analytes: (átomos de carbono, nome, [área rep1..3], [%Ai rep1..3])."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Dados {detector}"
    ws["A1"] = detector
    ws["A2"], ws["B2"] = "AMOSTRA:", cell_code
    ws.append([])
    ws.append(["ID", "Átomos Carbono", "Analito", "TR ( min)", "Replicatas", None, None, None, None, None, "%Ai = Ai/ATOTAL * 100"])
    ws.append([None, None, None, None, "Rep_1", "Rep_2", "Rep_3", "Media", "DSV%", None, "Rep_1", "Rep_2", "Rep_3", "Media", "DSV%"])
    for i, (carbon, name, areas, pcts) in enumerate(analytes, start=1):
        ws.append([i, carbon, name, 1.0 + i, *areas, sum(areas) / 3, 1.0, None, *pcts, sum(pcts) / 3, 1.0])
    ws.append([len(analytes) + 1, None, None, None, None, None, None, None, None, None, 0, 0, 0, 0, 0])
    wb.create_sheet("Entrada Calibração")
    return _xlsx(wb)


def gas_xlsx(code: str, initial_mass: float = 100.0, gas_mass: float | None = 1.0, gas: str = "Nitrogênio") -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados FID-TCD"
    rows = [
        {1: None, 3: "Valores a preencher", 8: "Amostra:", 9: code},
        {1: "DADOS EXPERIMENTO", 8: "REATOR", 9: "Reator utilizado", 11: "Reator X"},
        {1: "Experimento/Amostra:", 2: code, 9: "Diâmetro do reator", 11: 60, 12: "mm"},
        {1: "Reator utilizado", 2: "Reator X", 3: "← Selecione o reator utilizado"},
        {1: "Condições do teste"},
        {1: "Selecione o gás →", 2: gas},
        {1: "Massa de inicial de amostra", 2: initial_mass, 3: "g"},
        {1: "Pressão inicial", 2: 20, 3: "psi g", 9: "Massa de inicial de amostra", 11: initial_mass, 12: "g"},
        {1: "Massa de gás após pesagem", 2: 0.9, 3: "g", 9: "Massa de gás gerada", 11: gas_mass, 12: "g"},
        {1: "DADOS CROMATOGRAFIA", 9: "Fechamento do balanço de pressão (%)", 11: 1.5},
        {1: "Componente", 2: "Média", 3: "Média", 4: "Componente", 5: "Média"},
        {1: "Metano", 2: 6.0, 3: 60.0, 4: "H2", 5: 2.0},
        {1: "Etano", 2: 2.0, 3: 20.0, 4: "Metano", 5: 6.0},
        {1: "Propano", 2: 1.0, 3: 10.0, 4: "CO2", 5: 30.0},
        {1: "n-Butano", 2: 0.6, 3: 6.0},
        {1: "n-pentano", 2: 0.4, 3: 4.0},
    ]
    for r, cells in enumerate(rows, start=1):
        for col, value in cells.items():
            ws.cell(row=r, column=col, value=value)
    wb.create_sheet("Base reatores")
    return _xlsx(wb)


def pygcms_xlsx(sheets: dict[str, list[tuple[str, float]]]) -> bytes:
    """sheets: aba → [(identificação, área)]."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, peaks in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(["Análise de Rochas -Py-GC-MS "])
        ws.append(["Tempo de retenção", "m/z", "Area", "Area %", "Altura", "Altura%", "Area/Altura", "Identificação"])
        for i, (ident, area) in enumerate(peaks):
            ws.append([20 + i, 57, area, 1.0, area / 3, 1.0, 3.0, ident])
    return _xlsx(wb)


def literature_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Artigo"
    ws.append(["Table 6 — Synthetic compositional parameters of hydrous pyrolysis products."])
    ws.append([])
    ws.append(["Temp. (°C)", "Time (h)", "pristane/phytane"])
    ws.append([300, 72, 1.0])
    return _xlsx(wb)


def zip_of(files: dict[str, bytes]) -> bytes:
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, content in files.items():
            zf.writestr(name, content)
    return buf.getvalue()


def leco_ri_xlsx(rows: list[tuple], flat_copy: bool = False) -> bytes:
    """"Planilha de massas das amostras" do LECO, como a do laboratório: aba
    "Dados" com linha de grupos + linha de réplicas (1, 2, 3) e fórmulas nas
    colunas calculadas. rows: (data, amostra, [cadinho x3], [amostra x3], [após x3], obs).
    flat_copy: acrescenta a aba "Tabela" (layout plano) com os mesmos dados."""
    import datetime as dt

    import openpyxl
    from openpyxl.utils import get_column_letter as col

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(["Dados de massas das amostras"])
    groups = ["Massa do cadinho (g)", "Massa da amostra (g)", "Massa do cadinho com amostra (g)", "Massa após tratamento (g)", "Resíduo Insolúvel"]
    head = ["Data", "Amostra"]
    for g in groups:
        head += [g, None, None]
    ws.append(head + ["Observações"])
    ws.append([None, None] + ["1", "2", "3"] * len(groups))
    for r, (day, name, cad, amo, apos, obs) in enumerate(rows, start=4):
        line = [dt.datetime(2026, 1, day), name, *cad, *amo]
        line += [f"={col(3 + k)}{r}+{col(6 + k)}{r}" for k in range(3)]
        line += list(apos)
        line += [f"={col(9 + k)}{r}-{col(12 + k)}{r}" for k in range(3)]
        ws.append(line + [obs])
    if flat_copy:
        flat = wb.create_sheet("Tabela")
        flat.append(["Data", "Amostra", "Cadinho 1", "Cadinho 2", "Cadinho 3", "Amostra 1", "Amostra 2", "Amostra 3", "Após 1", "Após 2", "Após 3", "Observações"])
        for day, name, cad, amo, apos, obs in rows:
            flat.append([dt.datetime(2026, 1, day), name, *cad, *amo, *apos, obs])
    return _xlsx(wb)


def leco_ri_flat_xlsx(rows: list[tuple]) -> bytes:
    """Só o layout plano (aba "Tabela")."""
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabela"
    ws.append(["Data", "Amostra", "Cadinho 1", "Cadinho 2", "Cadinho 3", "Amostra 1", "Amostra 2", "Amostra 3", "Após 1", "Após 2", "Após 3", "Observações"])
    for day, name, cad, amo, apos, obs in rows:
        ws.append([dt.datetime(2026, 1, day), name, *cad, *amo, *apos, obs])
    return _xlsx(wb)


def gas_consolidated_xlsx(experiments: dict[str, list[tuple]], initial_masses: dict[str, float] | None = None) -> bytes:
    """Tabela consolidada de gás: aba "Detalhe" (uma linha por pico) + "Resumo".
    experiments: {código: [(componente agrupado, mol%, massa cromatografia g, massa pressão g)]}."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Tabela final"
    wb["Tabela final"].append(["Consolidado – Gás gerado (fórmulas da aba Detalhe)"])
    det = wb.create_sheet("Detalhe")
    det.append(["Amostra", "Linha na planilha original", "Componente (original)", "Componente (agrupado)", "MM (g/mol)",
                "mol% (final, sem N2)", "% mássico", "Massa gerada - cromatografia (g)", "Massa gerada - cromatografia (mg/g amostra)",
                "Massa gerada - pressão/pesagem (g)", "Massa gerada - pressão/pesagem (mg/g amostra)"])  # fmt: skip
    for code, comps in experiments.items():
        for i, (name, mol, crom, press) in enumerate(comps, start=11):
            det.append([code, i, name, name, None, mol, None, crom, None, press, None])
    det.append(["Fonte: nota sem código de experimento"])
    if initial_masses:
        res = wb.create_sheet("Resumo")
        res.append(["Resumo por amostra – massas geradas"])
        res.append([])
        res.append(["Parâmetro", *initial_masses])
        res.append(["Massa inicial de amostra (g)", *initial_masses.values()])
        res.append(["Pressão inicial (psi g)", *[25 for _ in initial_masses]])
    return _xlsx(wb)


# ---------------------------------------------------------------- tabaco (14/10/2026)
# Mesma estrutura dos arquivos do projeto de tabaco; nomes e números inventados.


def _tg_mass(temp: float) -> float:
    """Curva inventada: umidade (~100 °C, 8 %), pirólise principal (~320 °C,
    50 %), 2ª etapa (~480 °C, 10 %) → resíduo ~32 %."""
    import math

    def step(t0: float, width: float, amount: float) -> float:
        return amount / (1 + math.exp(-(temp - t0) / width))

    return 100 - step(100, 12, 8) - step(320, 25, 50) - step(480, 20, 10)


def tg_txt(sample_name: str, weight_mg: float = 8.0, unit: str = "%", encoding: str = "utf-16", t_end: float = 1000.0) -> bytes:
    """Exportação ASCII do TGA-50: utf-16 (variante nova, "File Name:" no topo)
    ou cp1252 ("[TA Ascii File]", os "-Corrected"). unit "%" ou "mg"."""
    head = []
    if encoding == "utf-16":
        head += [f"File Name:\tC:\\TAData\\Sintetico\\{sample_name}.tad", "", "[File Information]", f"File Name:\t{sample_name}.tad",
                 f"Sample Name:\t{sample_name}", "Acquisition Date\t26/05/04", "Acquisition Time\t10:20:30(-0300)", "Detector:\tTGA-50",
                 "Operator:\tSintetico", "Atmosphere:\tNitrogen", "Flow Rate:\t100[ml/min]", "Cell:\tPlatinum",
                 f"Sample Weight:\t{weight_mg:.3f}[mg]", "", "", "[Program]", "Temp Rate\tHold Temp\tHold Time", "[C/min ]\t[  C   ]\t[  min  ]",
                 "20.00\t1000.0\t0", "", ""]  # fmt: skip
    else:
        head += ["[TA Ascii File]", "TGA-50", "", f"File Name :\tC:\\TAData\\Sintetico\\{sample_name}-Corrected.txt", "", "[File Information]",
                 f"File Name:\t {sample_name} .tad", f"Sample Name:\t{sample_name} ", "Lot No:\t", "Acquire Date:\t26/05/04",
                 "Acquire Time:\t10:20:30(-0300)", "Detector Type:\tTGA-50", "Operator:\tSint\xe9tico", "Atmosphere:\tNitrogen",
                 "Gas Flow:\t100[ml/min]", f"Sample Weight:\t{weight_mg:.3f}[mg]", "", "", "[DataCorrectHistory]",
                 "Smoothing-TGA:\t26/05/05 10:00:00(-0300)", "\t100Points", "", ""]  # fmt: skip
    lines = [*head, "[Data]", "Time\tTemp\tTGA", f"sec\tC\t{unit}"]
    t0 = 25.0
    seconds = int((t_end - t0) / 20 * 60)
    for s in range(seconds + 1):
        temp = t0 + s * 20 / 60
        m = _tg_mass(temp) / _tg_mass(t0) * 100
        value = m if unit == "%" else m / 100 * weight_mg
        lines.append(f"{s:.2f}\t{temp:.2f}\t{value:.4f}")
    text = "\r\n".join(lines) + "\r\n\r\n"
    return text.encode("utf-16") if encoding == "utf-16" else text.encode("cp1252")


def eds_quantity_csv(elements: list[tuple[str, float, float]]) -> bytes:
    """"Spectrum NN Quantity analysis.csv" do AZtec: (nome em inglês, % atômica, % massa)."""
    lines = ["AV;20000", "PC;1e-09", "Total number of counts;1000000", "Mode;Quality", "Profile;Resolution", "Real time;40", "Live time;30", ""]
    lines.append("Element,Atomic %,Weight %")
    lines += [f"{name},{at:.2f}%,{wt:.2f}%" for name, at, wt in elements]
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


def eds_spectrum_csv() -> bytes:
    lines = ["AV;20000", "PC;1e-09", "Total number of counts;1000000", "", "Energy [eV];Counts"]
    lines += [f"{10 * i};{(i * 7) % 13}" for i in range(1, 50)]
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


def eds_msa() -> bytes:
    return b"#FORMAT      : EMSA/MAS Spectral Data File\r\n#VERSION     : 1.0\r\n#TITLE       : Spectrum\r\n#NPOINTS     : 3\r\n#SPECTRUM    :\r\n1, 2, 3\r\n#ENDOFDATA   :\r\n"


def leco_raw_csv(sets: list[tuple[str, int, float]], blank_height: float = 1.0, n_blanks: int = 3, seconds: int = 60) -> bytes:
    """Sinal bruto do Cornerstone (1 linha por segundo). sets: (nome, réplica,
    altura do pico de C); S baixo = 1/10 de C; S alto vazio às vezes."""
    out = ['"Set","Rep","Time","Carbon","Sulfur Low","Sulfur High","Furnace Power","Furnace Temperature"']
    clock = 0

    def run(name: str, rep: int, height: float) -> None:
        nonlocal clock
        for s in range(seconds):
            clock += 1
            hh, mm, ss = 10 + clock // 3600, (clock // 60) % 60, clock % 60
            c = height if 10 <= s < 20 else 0.0
            s_high = "" if s % 7 == 0 else f"{c / 100}"
            out.append(f'"{name}","{rep}","2026-05-04 {hh:02d}:{mm:02d}:{ss:02d}","{c}","{c / 10}","{s_high}","1200.0","1350.0"')

    for name, rep, height in sets:
        run(name, rep, height)
    for rep in range(1, n_blanks + 1):
        run("Blank", rep, blank_height)
    return ("﻿" + "\r\n".join(out) + "\r\n").encode("utf-8")


RE_MULTI_HEADER = ["Sample", "MINC (%)", "OI", "HI", "Q1 (mg/g)", "Q2 (mg/g)", "Q3 (mg/g)", "Q4 (mg/g)", "Q5 (mg/g)", "Tmax (°C)",
                   "TOC (%)", "Total HC Gas (%)", "Total HC Liquid (%)", "PC (%)", "RC (%)"]  # fmt: skip


def rockeval_multiheating_htm(rows: list[dict]) -> bytes:
    """Relatório "multiheating": SEM coluna Analysis; curveData na ordem das
    linhas, com o nome da corrida. rows: {"sample", "run", "toc", "q1", "curves": bool}."""
    header = ["\xa0" + h.replace(" ", "\xa0") + "\xa0" for h in RE_MULTI_HEADER]
    table = [[r["sample"], "0.40", "50", "300", str(r["q1"]), "10.0", "20.0", "5.0", "", "420", str(r["toc"]), "", "", "8.0", "40.0"] for r in rows]
    parts = ["<!DOCTYPE html><html><head><script>", 'var reportTitle = "Job report / sintetico";']
    parts.append("var dataTableHeader = " + json.dumps(header, ensure_ascii=False) + ";\n")
    parts.append("var dataTableArray = " + json.dumps(table) + ";\n")
    links = []
    for i, r in enumerate(rows):
        if r.get("curves", True):
            n = 30
            data = [list(range(n)), [200 + 10 * t for t in range(n)], [float(t % 5 + i) for t in range(n)], None, None, None, None, None, None, None]
            parts.append(f"var pyroData{i} = {json.dumps(data)};\n")
            links.append(f'["{r["run"]}",pyroData{i},null]')
        else:
            links.append(f'["{r["run"]}",null,null]')
    parts.append("var curveData = [" + ",\n".join(links) + "];\n")
    parts.append("</script></head><body></body></html>")
    return "".join(parts).encode("utf-8")
