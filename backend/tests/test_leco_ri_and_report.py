"""LECO - Resíduo Insolúvel (leitor novo) e o relatório em Excel do botão
"Exportar". Planilhas sintéticas, códigos genéricos."""

import io

import openpyxl
import pytest

from app.parsers import parse_file
from app.parsers.common import ParseError
from app.parsers.leco_ri import code_suggestion
from tests import synthetic as syn
from tests.conftest import IC, PESQ, confirm, upload

ROWS = [
    ("HP300H", "HP300NA", 0.10, 0.11, 0.12),
    ("HP320SE", "HP320NB", 0.20, 0.22, None),
    ("HP355 - SEM EXTRAÇÃO", "HP355NB", 0.30, 0.31, 0.32),
    ("HP300NA_E", "HP300NA", 0.05, 0.06, 0.07),
    ("Rocha teste", "Rocha virgem", 0.40, 0.41, 0.42),
    ("HP340H", "HP340NA", None, None, None),  # sem valores: ignorada
    (None, "HP360NA", None, None, None),  # só o experimento: ignorada
]


def test_parser_reads_by_header_and_suggests_codes():
    r = parse_file("Resíduo Insolúvel.xlsx", syn.leco_ri_xlsx(ROWS))
    assert r.technique == "leco_ri" and r.format_label == "LECO - Resíduo Insolúvel"
    by_name = {rec.raw_name: rec for rec in r.records}
    assert set(by_name) == {"HP300H", "HP320SE", "HP355 - SEM EXTRAÇÃO", "HP300NA_E", "Rocha teste"}
    assert [(v.parameter, v.value, v.replicate) for v in by_name["HP300H"].values] == [
        ("RI", 0.10, 1),
        ("RI", 0.11, 2),
        ("RI", 0.12, 3),
    ]
    assert len(by_name["HP320SE"].values) == 2  # réplica vazia não entra
    assert by_name["HP300H"].experiment.code == "HP300NA"
    assert by_name["Rocha teste"].experiment is None  # "Rocha virgem" não é código de experimento
    assert by_name["HP355 - SEM EXTRAÇÃO"].code_hint == "HP355SE"
    assert by_name["HP300NA_E"].code_hint == "HP300NAE"
    assert any("sem nenhuma réplica" in w for w in r.warnings)


def test_code_suggestions():
    assert code_suggestion("HP355NB.E") == "HP355NBE"
    assert code_suggestion("HP280NBE.1") is None  # alíquota: o código já é reconhecido
    assert code_suggestion("Rocha araripe 80mesh") is None


def test_sheet_without_values_is_error():
    with pytest.raises(ParseError):
        parse_file("x.xlsx", syn.leco_ri_xlsx([("HP300H", "HP300NA", None, None, None)]))


def test_import_uses_experiment_from_sheet(client, project):
    pid = project["id"]
    prev = upload(client, pid, {"Resíduo Insolúvel.xlsx": syn.leco_ri_xlsx(ROWS[:3])})
    confirm(client, pid, prev)
    samples = {s["code"]: s for s in client.get(f"projects/{pid}/samples", headers=PESQ).json()["samples"]}
    assert samples["HP300H"]["experiment_code"] == "HP300NA"
    assert samples["HP355SE"]["fraction"] == "SE"
    ri = samples["HP300H"]["values"]["leco_ri.RI"]
    assert ri["n"] == 3 and ri["mean"] == pytest.approx(0.11)


def _report(client, pid, body, who=PESQ):
    r = client.post(f"projects/{pid}/export/report", json=body, headers=who)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return openpyxl.load_workbook(io.BytesIO(r.content))


def test_report_has_summary_first_and_one_sheet_per_technique(client, project):
    pid = project["id"]
    confirm(client, pid, upload(client, pid, {"Resíduo Insolúvel.xlsx": syn.leco_ri_xlsx(ROWS[:3])}))
    confirm(
        client,
        pid,
        upload(client, pid, {"leco.csv": syn.leco_csv([{"name": "HP300H", "set_id": "S1", "reps": [(10.0, 1.0, 0.2, "Included"), (12.0, 1.2, 0.2, "Included")]}])}),
    )
    wb = _report(client, pid, {"techniques": [], "mode": "padrao"}, who=IC)  # IC do projeto também exporta
    assert wb.sheetnames[0] == "Resumo geral"
    assert "LECO - Resíduo Insolúvel" in wb.sheetnames and "LECO" in wb.sheetnames

    summary = wb["Resumo geral"]
    text = [str(c.value) for row in summary.iter_rows() for c in row if c.value is not None]
    assert any("Projeto Teste" in t for t in text)
    assert "Resumo por técnica" in text

    ri = wb["LECO - Resíduo Insolúvel"]
    cells = [[c.value for c in row] for row in ri.iter_rows()]
    # bloco "Todas as réplicas": uma linha por réplica (3 + 2 + 3)
    start = next(i for i, row in enumerate(cells) if row[0] == "Todas as réplicas")
    reps = [row for row in cells[start + 2 :] if row[0]]
    assert len(reps) == 8
    assert {row[0] for row in reps} == {"HP300H", "HP320SE", "HP355SE"}


def test_report_only_chosen_techniques_and_samples(client, project):
    pid = project["id"]
    confirm(client, pid, upload(client, pid, {"Resíduo Insolúvel.xlsx": syn.leco_ri_xlsx(ROWS[:3])}))
    samples = {s["code"]: s["id"] for s in client.get(f"projects/{pid}/samples", headers=PESQ).json()["samples"]}
    wb = _report(client, pid, {"techniques": ["leco_ri"], "sample_ids": [samples["HP300H"]], "mode": "todas"})
    assert wb.sheetnames == ["Resumo geral", "LECO - Resíduo Insolúvel"]
    codes = {c.value for row in wb["LECO - Resíduo Insolúvel"].iter_rows() for c in row[:1]}
    assert "HP300H" in codes and "HP320SE" not in codes


def test_report_rejects_unknown_technique(client, project):
    r = client.post(f"projects/{project['id']}/export/report", json={"techniques": ["xyz"]}, headers=PESQ)
    assert r.status_code == 400
