"""Relatório em Excel do projeto (botão "Exportar", ao lado de "Importar").

Pedido do mantenedor (09/10/2026): um documento organizado para apresentar —

* aba **"Resumo geral"** (a primeira): cabeçalho do projeto (quem exportou,
  quando, quais resultados entraram), um quadro por técnica (amostras,
  medições, réplicas, período) e a tabela por amostra com média e desvio dos
  parâmetros principais de cada técnica escolhida;
* **uma aba por técnica** escolhida: as médias por amostra de todos os
  parâmetros e, embaixo, **todas as réplicas** (uma linha por réplica de
  cada medição, com arquivo, data e validade).

A pessoa escolhe as técnicas, as amostras e o modo de validade (padrão /
só válidas / todas — o mesmo das telas). Valores = os mesmos das telas
(`results.load`, inclusive a massa de gás editada).
"""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlmodel import Session, select

from app.db.models import FractionType, Project, StoredFile
from app.services import results
from app.services.catalog import TECHNIQUE_ORDER, TECHNIQUES

MODE_LABELS = {
    "padrao": "Padrão — válidas e pendentes (inválidas de fora)",
    "validas": "Só amostras marcadas como válidas",
    "todas": "Todas, inclusive as inválidas",
}
VALIDITY = {True: "válida", False: "inválida", None: "pendente"}

# Cores do Horun (azul-escuro) com tons claros para leitura em papel.
NAVY = "15216F"
LIGHT = "E8EBF7"
ZEBRA = "F6F7FB"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="C9CEDF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name="Calibri", size=12, bold=True, color=NAVY)
HEAD_FONT = Font(name="Calibri", size=10, bold=True, color=WHITE)
GROUP_FONT = Font(name="Calibri", size=10, bold=True, color=NAVY)
BODY_FONT = Font(name="Calibri", size=10)
MUTED_FONT = Font(name="Calibri", size=9, italic=True, color="666666")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
GROUP_FILL = PatternFill("solid", fgColor=LIGHT)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)

SAMPLE_COLS = ["Amostra", "Fração", "Temperatura (°C)", "Experimento", "Atmosfera", "Réplica do exp.", "Validade"]


def _num_format(values: list[float]) -> str:
    big = max((abs(v) for v in values if v is not None), default=0)
    if big >= 100:
        return "0.0"
    if big >= 1:
        return "0.00"
    return "0.0000"


def _param_label(p) -> str:
    return f"{p.label} ({p.unit})" if p.unit else p.label


def _safe_sheet(title: str, used: set[str]) -> str:
    clean = "".join("-" if ch in '[]:*?/\\' else ch for ch in title)[:31] or "Aba"
    name, n = clean, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name, n = clean[: 31 - len(suffix)] + suffix, n + 1
    used.add(name.lower())
    return name


def _head(ws: Worksheet, row: int, col: int, text: str, width: int = 1, group: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = GROUP_FONT if group else HEAD_FONT
    cell.fill = GROUP_FILL if group else HEAD_FILL
    cell.alignment = CENTER
    cell.border = BORDER
    if width > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width - 1)
        for c in range(col + 1, col + width):
            ws.cell(row=row, column=c).border = BORDER
            ws.cell(row=row, column=c).fill = GROUP_FILL if group else HEAD_FILL


def _body(ws: Worksheet, row: int, col: int, value, zebra: bool, fmt: str | None = None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = BODY_FONT
    cell.border = BORDER
    cell.alignment = LEFT if isinstance(value, str) else Alignment(horizontal="right", vertical="center")
    if zebra:
        cell.fill = ZEBRA_FILL
    if fmt and isinstance(value, (int, float)):
        cell.number_format = fmt
    return cell


def _widths(ws: Worksheet, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def build(
    session: Session,
    project: Project,
    techniques: list[str],
    sample_ids: list[int],
    mode: str,
    exported_by: str,
) -> bytes:
    data = results.load(session, project.id, sample_ids or None)
    fractions = {f.code: f.label for f in session.exec(select(FractionType))}
    files = {f.id: f.filename for f in session.exec(select(StoredFile).where(StoredFile.project_id == project.id))}

    samples = {sid: s for sid, s in data.samples.items() if results._sample_ok(s, mode)}
    analyses = {
        aid: a
        for aid, a in data.analyses.items()
        if a.sample_id in samples and results._analysis_ok(a, mode) and a.technique in techniques
    }
    values_by_analysis: dict[int, list] = defaultdict(list)
    for v in data.values:
        if v.analysis_id in analyses:
            values_by_analysis[v.analysis_id].append(v)
    chosen = [t for t in TECHNIQUE_ORDER if t in techniques and any(a.technique == t for a in analyses.values())]

    def order(s):
        return (s.kind != "sample", s.temperature_c is None, s.temperature_c or 0, s.code)

    with_data = {a.sample_id for a in analyses.values()}
    ordered = [s for s in sorted(samples.values(), key=order) if s.id in with_data]

    # (técnica, amostra, parâmetro) -> valores
    pooled: dict[tuple[str, int, str], list[float]] = defaultdict(list)
    for aid, vals in values_by_analysis.items():
        a = analyses[aid]
        for v in vals:
            pooled[(a.technique, a.sample_id, v.parameter)].append(v.value)

    def sample_cells(s) -> list:
        row = results.sample_row(s, data.experiments)
        return [
            s.code,
            fractions.get(s.fraction, s.fraction),
            s.temperature_c,
            row["experiment_code"] or "",
            row["atmosphere"] or "",
            row["replicate_letter"] or "",
            VALIDITY[s.valid],
        ]

    wb = Workbook()
    used: set[str] = set()
    summary = wb.active
    summary.title = _safe_sheet("Resumo geral", used)
    _summary_sheet(summary, project, chosen, ordered, analyses, values_by_analysis, pooled, sample_cells, mode, exported_by)

    for tech in chosen:
        ws = wb.create_sheet(_safe_sheet(TECHNIQUES[tech]["label"], used))
        _technique_sheet(
            ws, tech, ordered, analyses, values_by_analysis, pooled, sample_cells, files, data.samples, data.experiments, fractions
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary_sheet(ws, project, chosen, ordered, analyses, values_by_analysis, pooled, sample_cells, mode, exported_by) -> None:
    # horário de Brasília (sem horário de verão desde 2019; não depende do tzdata)
    now = datetime.now(timezone(timedelta(hours=-3))).strftime("%d/%m/%Y %H:%M")
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Relatório de resultados — {project.name}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Horun · Resultados — Laboratório NQTR, Instituto de Química (UFRJ)"
    ws["A2"].font = MUTED_FONT

    n_meas = len(analyses)
    n_reps = sum(len({v.replicate for v in vals}) or 1 for vals in values_by_analysis.values())
    info = [
        ("Projeto", project.name),
        ("Descrição", project.description or "—"),
        ("Exportado em", now),
        ("Exportado por", exported_by),
        ("Resultados considerados", MODE_LABELS.get(mode, mode)),
        ("Amostras", len(ordered)),
        ("Medições", n_meas),
        ("Técnicas", ", ".join(TECHNIQUES[t]["label"] for t in chosen) or "—"),
    ]
    r = 4
    for label, value in info:
        k = ws.cell(row=r, column=1, value=label)
        k.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        v = ws.cell(row=r, column=2, value=value)
        v.font = BODY_FONT
        v.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        r += 1

    # quadro por técnica
    r += 1
    ws.cell(row=r, column=1, value="Resumo por técnica").font = SUBTITLE_FONT
    r += 1
    for i, title in enumerate(["Técnica", "Instrumento", "Amostras", "Medições", "Réplicas", "Primeira análise", "Última análise"], start=1):
        _head(ws, r, i, title)
    r += 1
    for idx, tech in enumerate(chosen):
        tech_analyses = [a for a in analyses.values() if a.technique == tech]
        dates = sorted(a.analyzed_at for a in tech_analyses if a.analyzed_at)
        reps = sum(len({v.replicate for v in values_by_analysis[a.id]}) or 1 for a in tech_analyses)
        cells = [
            TECHNIQUES[tech]["label"],
            TECHNIQUES[tech].get("instrument", ""),
            len({a.sample_id for a in tech_analyses}),
            len(tech_analyses),
            reps,
            dates[0][:10] if dates else "—",
            dates[-1][:10] if dates else "—",
        ]
        for c, value in enumerate(cells, start=1):
            _body(ws, r, c, value, idx % 2 == 1, "0")
        r += 1
    _body(ws, r, 1, "Total", False).font = Font(name="Calibri", size=10, bold=True)
    for c, value in ((3, len(ordered)), (4, n_meas), (5, n_reps)):
        _body(ws, r, c, value, False, "0").font = Font(name="Calibri", size=10, bold=True)
    r += 2

    # tabela por amostra — parâmetros principais
    ws.cell(row=r, column=1, value="Resultados por amostra — média e desvio padrão (DP) dos parâmetros principais").font = SUBTITLE_FONT
    r += 1
    group_row, head_row = r, r + 1
    for i, title in enumerate(SAMPLE_COLS, start=1):
        _head(ws, group_row, i, "Amostra" if i == 1 else "", 1, group=True)
        _head(ws, head_row, i, title)
    ws.merge_cells(start_row=group_row, start_column=1, end_row=group_row, end_column=len(SAMPLE_COLS))
    col = len(SAMPLE_COLS) + 1
    columns: list[tuple[str, str, str]] = []  # (técnica, parâmetro, formato)
    for tech in chosen:
        params = [p for p in TECHNIQUES[tech]["params"] if p.main] or TECHNIQUES[tech]["params"][:3]
        params = [p for p in params if any(pooled.get((tech, s.id, p.key)) for s in ordered)]
        if not params:
            continue
        _head(ws, group_row, col, TECHNIQUES[tech]["label"], 2 * len(params), group=True)
        for p in params:
            vals = [x for s in ordered for x in pooled.get((tech, s.id, p.key), [])]
            fmt = _num_format(vals)
            _head(ws, head_row, col, _param_label(p))
            _head(ws, head_row, col + 1, "DP")
            columns.append((tech, p.key, fmt))
            col += 2
    r = head_row + 1
    first_data = r
    for idx, s in enumerate(ordered):
        zebra = idx % 2 == 1
        for c, value in enumerate(sample_cells(s), start=1):
            _body(ws, r, c, value if value is not None else "", zebra, "0")
        c = len(SAMPLE_COLS) + 1
        for tech, key, fmt in columns:
            stats = results.mean_sd(pooled.get((tech, s.id, key), []))
            _body(ws, r, c, stats["mean"] if stats["mean"] is not None else "", zebra, fmt)
            sd = _body(ws, r, c + 1, stats["sd"] if stats["sd"] is not None else "", zebra, fmt)
            sd.font = Font(name="Calibri", size=9, color="555555")
            c += 2
        r += 1
    if ordered:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(max(col - 1, len(SAMPLE_COLS)))}{r - 1}"
    ws.freeze_panes = ws.cell(row=first_data, column=2)

    r += 1
    notes = [
        "Média e DP calculados com todas as réplicas de todas as medições da amostra (DP amostral, n−1). "
        "Detalhe de cada técnica — todos os parâmetros e cada réplica — nas abas seguintes.",
        "Validade: válida = conferida; pendente = ainda não revisada; inválida = descartada "
        "(só entra quando o relatório pede “todas”).",
        "Massa de gás: vale a editada em Condições experimentais, quando houver.",
    ]
    for note in notes:
        ws.cell(row=r, column=1, value=note).font = MUTED_FONT
        r += 1

    _widths(ws, {1: 24, 2: 22, 3: 13, 4: 14, 5: 13, 6: 12, 7: 12})
    for c in range(len(SAMPLE_COLS) + 1, col):
        ws.column_dimensions[get_column_letter(c)].width = 13 if (c - len(SAMPLE_COLS)) % 2 == 1 else 9
    ws.row_dimensions[head_row].height = 30
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{group_row}:{head_row}"


def _technique_sheet(ws, tech, ordered, analyses, values_by_analysis, pooled, sample_cells, files, all_samples, experiments, fractions) -> None:
    info = TECHNIQUES[tech]
    params = [p for p in info["params"] if any(pooled.get((tech, s.id, p.key)) for s in ordered)]
    ws.sheet_view.showGridLines = False
    ws["A1"] = info["label"]
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Instrumento: {info.get('instrument', '')}"
    ws["A2"].font = MUTED_FONT
    fmts = {p.key: _num_format([x for s in ordered for x in pooled.get((tech, s.id, p.key), [])]) for p in params}

    # 1) médias por amostra
    r = 4
    ws.cell(row=r, column=1, value="Médias por amostra").font = SUBTITLE_FONT
    r += 1
    for i, title in enumerate(SAMPLE_COLS, start=1):
        _head(ws, r + 1, i, title)
        _head(ws, r, i, "", group=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(SAMPLE_COLS))
    col = len(SAMPLE_COLS) + 1
    for p in params:
        _head(ws, r, col, _param_label(p), 3, group=True)
        for j, title in enumerate(("Média", "DP", "n")):
            _head(ws, r + 1, col + j, title)
        col += 3
    last_col = col - 1
    r += 2
    rows_here = [s for s in ordered if any(a.sample_id == s.id and a.technique == tech for a in analyses.values())]
    for idx, s in enumerate(rows_here):
        zebra = idx % 2 == 1
        for c, value in enumerate(sample_cells(s), start=1):
            _body(ws, r, c, value if value is not None else "", zebra, "0")
        c = len(SAMPLE_COLS) + 1
        for p in params:
            stats = results.mean_sd(pooled.get((tech, s.id, p.key), []))
            _body(ws, r, c, stats["mean"] if stats["mean"] is not None else "", zebra, fmts[p.key])
            _body(ws, r, c + 1, stats["sd"] if stats["sd"] is not None else "", zebra, fmts[p.key])
            _body(ws, r, c + 2, stats["n"] or "", zebra, "0")
            c += 3
        r += 1

    # 2) todas as réplicas
    r += 2
    ws.cell(row=r, column=1, value="Todas as réplicas").font = SUBTITLE_FONT
    r += 1
    rep_cols = SAMPLE_COLS[:6] + [
        "Nome no arquivo",
        "Alíquota",
        "Réplica de análise",
        "Réplica",
        "Data da análise",
        "Arquivo",
        "Validade da medição",
    ]
    for i, title in enumerate(rep_cols, start=1):
        _head(ws, r, i, title)
    for j, p in enumerate(params):
        _head(ws, r, len(rep_cols) + 1 + j, _param_label(p))
    head_row = r
    r += 1
    tech_analyses = sorted(
        (a for a in analyses.values() if a.technique == tech),
        key=lambda a: ([s.id for s in ordered].index(a.sample_id), a.aliquot or 0, a.replicate or 0, a.analyzed_at, a.id),
    )
    idx = 0
    for a in tech_analyses:
        s = all_samples[a.sample_id]
        by_rep: dict[int | None, dict[str, float]] = defaultdict(dict)
        for v in values_by_analysis[a.id]:
            by_rep[v.replicate][v.parameter] = v.value
        general = by_rep.pop(None, {})
        reps = sorted(by_rep) or [None]
        for rep in reps:
            zebra = idx % 2 == 1
            vals = {**general, **by_rep.get(rep, {})} if rep is not None else general
            base = sample_cells(s)[:6] + [
                a.source_name or "",
                a.aliquot if a.aliquot is not None else "",
                a.replicate if a.replicate is not None else "",
                rep if rep is not None else "—",
                a.analyzed_at or "",
                files.get(a.source_file_id, "") if a.source_file_id else "",
                VALIDITY[a.valid],
            ]
            for c, value in enumerate(base, start=1):
                _body(ws, r, c, value if value is not None else "", zebra, "0")
            for j, p in enumerate(params):
                value = vals.get(p.key)
                _body(ws, r, len(rep_cols) + 1 + j, value if value is not None else "", zebra, fmts[p.key])
            r += 1
            idx += 1
    if tech_analyses:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(rep_cols) + len(params))}{r - 1}"
    ws.freeze_panes = ws.cell(row=7, column=2)

    widths = {1: 22, 2: 20, 3: 12, 4: 13, 5: 12, 6: 10, 7: 20, 8: 9, 9: 10, 10: 9, 11: 16, 12: 28, 13: 12}
    _widths(ws, widths)
    for c in range(len(SAMPLE_COLS) + 1, max(last_col, len(rep_cols) + len(params)) + 1):
        if c not in widths:
            ws.column_dimensions[get_column_letter(c)].width = 12
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
