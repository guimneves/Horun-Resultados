"""Detecção de formato + leitura (ESPECIFICACAO.md, seção 3).

O formato é detectado pelo CONTEÚDO (cabeçalhos/estrutura), não só pela
extensão. `parse_file` devolve um `ParseResult` ou levanta:
* `NotResultsFile` — arquivo reconhecido que não é de resultados (fica "ignorado");
* `ParseError` — arquivo não reconhecido ou quebrado (fica "erro").
Os .zip são abertos antes, em services/importer.py.
"""

from __future__ import annotations

import io

from app.parsers import chnso, gas_consolidated, gc, leco, leco_raw, leco_ri, mev_eds, pygcms, rockeval, shimadzu, tg
from app.parsers.common import NotResultsFile, ParseError, ParseResult
from app.services.derived import add_derived

__all__ = ["NotResultsFile", "ParseError", "ParseResult", "parse_file"]


def _pdf_first_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return reader.pages[0].extract_text() or "" if reader.pages else ""
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Não foi possível ler o PDF ({exc.__class__.__name__}).") from exc


def _in_eds_folder(path_hint: str) -> bool:
    return any(mev_eds.is_eds_dir(p) or p.strip().upper() == "MEV" for p in path_hint.replace("\\", "/").split("/"))


# PNG, JPEG, TIFF (II/MM), BMP, GIF — micrografias e figuras do MEV
_IMAGE_MAGIC = (b"\x89PNG", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*", b"BM", b"GIF8")


def _parse_pdf(content: bytes, filename: str, path_hint: str = "") -> ParseResult:
    text = _pdf_first_text(content)
    if chnso.is_eurovector(" ".join(text.split())):
        return chnso.parse(content, filename)
    leco.reject_if_diagnostic_pdf(text)
    if _in_eds_folder(path_hint):
        raise NotResultsFile("PDF do relatório do MEV-EDS (figuras e espectro) — os números vêm do \"Quantity analysis.csv\".")
    raise ParseError(
        "PDF não reconhecido. Do CHNSO, importe o \"Results Summary for Element %\"; "
        "do LECO, o CSV exportado do Cornerstone."
    )


def _parse_xlsx(content: bytes, filename: str, path_hint: str) -> ParseResult:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ParseError("Leitor de planilhas (openpyxl) não instalado no servidor.") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Não foi possível abrir a planilha ({exc.__class__.__name__}).") from exc
    try:
        names = wb.sheetnames
        upper = {n.strip().upper(): n for n in names}
        if "DADOS FID-TCD" in upper:
            return gc.parse_gas_balance(wb, upper["DADOS FID-TCD"], filename, path_hint)
        if "DADOS FID" in upper:
            return gc.parse_detector(wb, upper["DADOS FID"], filename, path_hint)
        if "DADOS TCD" in upper:
            return gc.parse_detector(wb, upper["DADOS TCD"], filename, path_hint)
        if pygcms.is_pygcms_workbook(wb):
            return pygcms.parse_workbook(wb, filename)
        if gas_consolidated.is_consolidated_gas_workbook(wb):
            return gas_consolidated.parse_workbook(wb, filename)
        if leco_ri.is_ri_workbook(wb):
            return leco_ri.parse_workbook(wb, filename)
        first_rows = " ".join(
            str(c) for ws in wb.worksheets[:1] for r in ws.iter_rows(min_row=1, max_row=6, values_only=True) for c in r if c
        )
        if "Table" in first_rows and ("Temp" in first_rows or "pyrolysis" in first_rows):
            raise NotResultsFile(
                "Esta planilha é uma tabela de literatura (dados publicados). A importação de tabelas de "
                "referência para sobrepor nos gráficos ainda não está disponível nesta versão."
            )
        raise ParseError(
            "Planilha não reconhecida. Formatos aceitos: \"Dados FID\", \"Dados TCD\", "
            "\"Planilha cálculo gás\" (aba \"Dados FID-TCD\"), Py-GC-MS (uma aba por amostra) e "
            "LECO - Resíduo Insolúvel (planilha de massas das amostras) e a tabela consolidada de gás (aba \"Detalhe\")."
        )
    finally:
        wb.close()


def parse_file(filename: str, content: bytes, path_hint: str = "") -> ParseResult:
    base = filename.replace("\\", "/").rsplit("/", 1)[-1]
    if base.startswith("~$"):
        raise NotResultsFile("Arquivo temporário do Excel (~$...), ignorado.")
    if not content:
        raise ParseError("Arquivo vazio.")
    if content[:5] == b"%PDF-":
        result = _parse_pdf(content, base, path_hint)
    elif content[:2] == b"PK":
        result = _parse_xlsx(content, base, path_hint)
    elif shimadzu.is_shimadzu(content):  # .gcd (GC-FID) / .qgd (GC-MS) — arquivo composto do Windows
        result = shimadzu.parse(content, base)
    elif content.startswith(_IMAGE_MAGIC):
        raise NotResultsFile("Imagem (micrografia/figura) — não é arquivo de resultados.")
    elif tg.is_binary_tad(content):
        result = tg.parse(content, base)  # levanta NotResultsFile com o pedido de exportar em ASCII
    elif content[:2] in (b"\xff\xfe", b"\xfe\xff"):  # texto UTF-16 (exportação do TA-60)
        if not tg.looks_like_tg(tg.decode(content)):
            raise ParseError("Arquivo de texto UTF-16 não reconhecido (esperado: exportação ASCII do TG).")
        result = tg.parse(content, base)
    else:
        head = content[:200000].decode("utf-8-sig", errors="replace")
        if mev_eds.looks_like_eds_spectrum(head):
            mev_eds.reject_spectrum()
        if mev_eds.looks_like_eds_quantity(head):
            result = mev_eds.parse(content, base, path_hint)
        elif tg.looks_like_tg(tg.decode(content[:20000])):
            result = tg.parse(content, base)
        elif leco_raw.looks_like_leco_raw(leco.decode(content[:20000])):
            result = leco_raw.parse(content, base)
        elif rockeval.looks_like_rockeval(head) or (
            "<html" in head.lower() and rockeval.looks_like_rockeval(content.decode("utf-8-sig", errors="replace"))
        ):
            result = rockeval.parse(content, base)
        elif leco.looks_like_leco_csv(leco.decode(content[:20000]) if content else ""):
            result = leco.parse(content, base)
        elif "<html" in head.lower():
            raise ParseError("Este HTML não é um relatório do Rock-Eval (GeoWorks \"Job report\").")
        else:
            raise ParseError(
                "Formato não reconhecido. Aceitos: PDF do CHNSO (EuroVector), CSV do LECO (resultados ou sinal "
                "bruto), relatório .htm do Rock-Eval, planilhas de GC-FID/TCD, de cálculo de gás e de Py-GC-MS, "
                "TG (ASCII do TGA-50), MEV-EDS (\"Quantity analysis.csv\") e GC-FID/GC-MS da Shimadzu (.gcd/.qgd) — ou um .zip com eles."
            )
    for record in result.records:
        add_derived(record)
    return result
